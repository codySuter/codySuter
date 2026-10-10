"""Builds the synthetic Eagle exports the Go tests read.

The layout copies a real Eagle location export (Sheet1, a blank first
column of "o" markers, "Location N" / "Current Location N" column pairs,
multi-line headings) but every SKU and description is made up.

    pip install xlwt openpyxl
    python3 testdata/make_fixtures.py
"""
import csv
import os
import random

import openpyxl
import xlwt

HERE = os.path.dirname(os.path.abspath(__file__))
HEAD = [" ", "SKU", "Description", "Location\nCodes", "Current\nLocation\nCodes"]
for n in range(2, 7):
    HEAD += [f"Location\n{n}", f"Current\nLocation\n{n}"]

# (sku, desc, loc1..loc6) — the hand-picked cases the tests assert on.
ROWS = [
    ("70013", 'ACE SHOVEL SQ PT 27"', "12R02", "", "6", "", "", ""),
    ("70018", "ACE SHOVEL RND PT", "12R02", "", "5", "USTOR", "", ""),
    ("779600", "TEST RAKE 16 TINE", "12R07", "", "100", "12R06", "", ""),
    ("3008391", "TEST HOSE 50FT", "14L05", "", "4", "12R01", "", ""),
    ("6209563", "TEST FLAG ONLY", "", "12R06", "", "", "", ""),
    ("6707640", "TEST MARKDOWN ITEM", "12R08", "MDONE", "", "", "", ""),
    ("7000137D", "TEST PRUNER — BYPASS™", "12R05", "", "3", "", "", ""),
    ("9087035", "TEST OVERSTOCK ONLY", "", "", "", "12R04", "", ""),
    ("5555555", "OTHER AISLE ITEM", "14L02", "", "12", "107", "", ""),
]


def numeric(v):
    """Eagle writes some cells as numbers — mimic that for capacities."""
    return int(v) if v.isdigit() else v


def write_xls(path, rows, numeric_cells=True):
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("Sheet1")
    for c, h in enumerate(HEAD):
        ws.write(0, c, h)
    for r, row in enumerate(rows, start=1):
        sku, desc, *locs = row
        ws.write(r, 0, "o")
        ws.write(r, 1, sku)
        ws.write(r, 2, desc)
        for i, v in enumerate(locs):
            col = 4 + 2 * i
            if v != "":
                ws.write(r, col, numeric(v) if numeric_cells else v)
    wb.save(path)


def big_rows():
    """~1200 SKUs with unique descriptions, so the string table spans many
    CONTINUE records (and some strings, with a wide character, cross one)."""
    rnd = random.Random(7)
    out = list(ROWS)
    sides = ["R", "L"]
    for i in range(1200):
        aisle = rnd.randint(1, 20)
        loc1 = f"{aisle}{rnd.choice(sides)}{rnd.randint(1, 12):02d}"
        desc = f"SYNTH ITEM {i:05d} {'—' if i % 7 == 0 else '-'} {rnd.randint(1, 99)}PK"
        over = rnd.choice(["", "", "USTOR", "107", "42", f"{aisle}R{rnd.randint(1, 9):02d}"])
        out.append((str(8000000 + i), desc, loc1, "", str(rnd.randint(1, 24)), over, "", ""))
    return out


def write_compass(path):
    """Like a Compass item export: "Item Number" for the SKU, plain
    "Location" … "Location 6" columns, the description last, every cell
    text (Compass writes numbers like a capacity of 6 as "6")."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Item Number", "Location"] + [f"Location {n}" for n in range(2, 7)] + ["Item Description"])
    for sku, desc, *locs in ROWS:
        ws.append([sku] + list(locs) + [desc])
    wb.save(path)


def main():
    write_compass(os.path.join(HERE, "compass-export.xlsx"))
    write_xls(os.path.join(HERE, "eagle-sample.xls"), ROWS)
    write_xls(os.path.join(HERE, "eagle-big.xls"), big_rows())

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Locations"
    ws.append(HEAD)
    for sku, desc, *locs in ROWS:
        line = ["o", sku, desc]
        for v in locs:
            line += [None, numeric(v) if v else None]
        ws.append(line)
    wb.save(os.path.join(HERE, "eagle-sample.xlsx"))

    with open(os.path.join(HERE, "eagle-import.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(["SKU"] + [f"Location {n}" for n in range(1, 7)])
        for sku, _desc, *locs in ROWS:
            w.writerow([sku] + list(locs))


if __name__ == "__main__":
    main()
